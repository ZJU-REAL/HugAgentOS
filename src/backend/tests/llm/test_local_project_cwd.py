"""Local project paths agree across the public tool/runner/preview boundaries."""

import json
from types import SimpleNamespace

import pytest

from tests.llm.test_local_project_pin import local_project, pin_tool
from core.llm.tool_collector import ToolCollector


def payload(result):
    block = result.content[0]
    return json.loads(block["text"] if isinstance(block, dict) else block.text)


async def test_relative_write_lands_in_project_and_pin_keeps_original(local_project, monkeypatch):
    from core.sandbox.script_runner_provider import ScriptRunnerProvider

    monkeypatch.setattr("core.sandbox.get_sandbox_provider", ScriptRunnerProvider)
    from core.llm.tools.write_tool import register_write
    from core.llm.tools._state import ReadStateTracker
    from core.llm.tool_permissions import (
        CURRENT_PERMISSION_TICKET,
        PermissionTicket,
        PermissionIntent,
    )

    source, scope = local_project
    destination = source.parent / "交付 results" / "report.html"
    collector = ToolCollector()
    register_write(
        collector,
        chat_id="pin-chat",
        sandbox_session_id="child-session",
        user_id="pin-user",
        state=ReadStateTracker(),
        scope=scope,
    )
    token = CURRENT_PERMISSION_TICKET.set(
        PermissionTicket(
            "Write",
            "write",
            "",
            "test-write",
            tuple(
                PermissionIntent("local_path", mode, str(destination), "write report")
                for mode in ("read", "write")
            ),
        )
    )
    try:
        result = payload(
            await collector.get_tool("Write")._func(
                file_path="交付 results/report.html", content="<h1>Project original</h1>"
            )
        )
    finally:
        CURRENT_PERMISSION_TICKET.reset(token)
    assert result.get("ok"), result
    assert destination.read_text() == "<h1>Project original</h1>"
    pinned = payload(
        await pin_tool(scope, "child-session")(file_paths=["交付 results/report.html"])
    )
    assert pinned.get("ok"), pinned


@pytest.fixture
async def local_tools(local_project, tmp_path, monkeypatch):
    import asyncio
    import socket
    import uvicorn
    from core.sandbox.script_runner_provider import ScriptRunnerProvider
    from services.script_runner_service import server
    from core.llm.tools import _paths
    from core.llm.tools._state import ReadStateTracker
    from core.llm.tools.write_tool import register_write
    from core.llm.tools.read_tool import register_read
    from core.llm.tools.edit_tool import register_edit
    from core.llm.tools.glob_tool import register_glob
    from core.llm.tools.grep_tool import register_grep
    from core.llm.tools.sandbox_tool import register_bash
    from core.llm.tools.pin_tool import register_pin_to_workspace
    from core.llm.tool_permissions import (
        ToolPermissionRegistry,
        ToolPermissionService,
        PermissionRuntime,
        CURRENT_PERMISSION_TICKET,
    )

    source, scope = local_project
    scratch_root = tmp_path / "session-storage"
    scratch_root.mkdir()
    scratch = str(scratch_root)
    monkeypatch.setattr(server, "WORKSPACE_ROOT", scratch)
    monkeypatch.setattr("core.sandbox._common.WORKSPACE", scratch)
    monkeypatch.setattr(_paths, "WORKSPACE_ROOT", scratch)
    monkeypatch.setattr(server, "_AUTH_TOKEN", "")
    monkeypatch.setenv("DEPLOY_PROFILE", "local")
    monkeypatch.setenv("SANDBOX_TOOLS_ENABLED", "true")
    sock = socket.socket()
    sock.bind(("127.0.0.1", 0))
    http_server = uvicorn.Server(uvicorn.Config(server.app, log_level="error", lifespan="off"))
    task = asyncio.create_task(http_server.serve(sockets=[sock]))
    for _ in range(100):
        if http_server.started:
            break
        await asyncio.sleep(0.01)
    assert http_server.started
    monkeypatch.setattr(
        ScriptRunnerProvider,
        "_base_url",
        property(lambda self: f"http://127.0.0.1:{sock.getsockname()[1]}"),
    )
    provider = ScriptRunnerProvider()
    monkeypatch.setattr("core.sandbox.get_sandbox_provider", lambda: provider)
    collector = ToolCollector()
    state = ReadStateTracker()
    common = dict(
        chat_id="pin-chat", sandbox_session_id="cwd-child", user_id="pin-user", scope=scope
    )
    for register in (register_read, register_write, register_edit):
        register(collector, state=state, **common)
    for register in (register_glob, register_grep):
        register(collector, **common)
    register_bash(collector, loader=None, loaded_skill_ids=set(), **common)
    register_pin_to_workspace(collector, scope=scope, sandbox_session_id="cwd-child")
    registry = ToolPermissionRegistry()
    for name, spec in collector.permission_specs.items():
        registry.register(name, spec, source="integration")
    service = ToolPermissionService(
        registry,
        PermissionRuntime(
            chat_id="pin-chat",
            sandbox_session_id="cwd-child",
            user_id="pin-user",
            interactive=False,
            approval_available=False,
            approval_mode="full",
            project_scope=scope,
        ),
    )

    async def invoke(name, **args):
        outcome = await service.authorize(SimpleNamespace(name=name, id="e2e-call", input=args))
        if not outcome.proceed:
            return outcome.payload
        token = CURRENT_PERMISSION_TICKET.set(outcome.ticket)
        try:
            return payload(await collector.get_tool(name)._func(**args))
        finally:
            CURRENT_PERMISSION_TICKET.reset(token)

    invoke.service = service
    try:
        yield invoke, source.parent, scope, provider
    finally:
        await server.process_sessions.close_all()
        http_server.should_exit = True
        await asyncio.wait_for(task, 5)
        sock.close()


async def test_bash_files_pin_and_authenticated_preview_end_to_end(
    local_tools, db_session, monkeypatch
):
    import base64
    import io
    import shlex
    import sys
    import zipfile
    import httpx
    from dataclasses import replace
    from fastapi import FastAPI
    from openpyxl import load_workbook
    from api.routes import files
    from core.config.settings import settings
    from core.db.engine import get_db
    from core.db.models import UserShadow

    invoke, root, scope, provider = local_tools
    generated = root / "交付 results"
    script = """from pathlib import Path
from openpyxl import Workbook
from zipfile import ZipFile
print(Path.cwd())
p=Path('交付 results'); p.mkdir()
(p/'report.html').write_text('<!doctype html><meta charset="utf-8"><h1>产业链 L4-L5</h1>')
w=Workbook(); w.active.append(['L4', 'L5']); w.active.append(['材料', '高纯材料']); w.save(p/'nodes.xlsx')
with ZipFile(p/'delivery.zip', 'w') as z:
    z.write(p/'report.html', 'report.html'); z.write(p/'nodes.xlsx', 'nodes.xlsx')
"""
    result = await invoke(
        "bash", command=shlex.quote(sys.executable) + " -c " + shlex.quote(script)
    )
    assert result.get("exit_code") == 0, result
    assert result["stdout"].strip() == str(root)
    assert not list(root.glob(".__process*"))
    for name in ("report.html", "nodes.xlsx", "delivery.zip"):
        assert (generated / name).is_file()

    read = await invoke("Read", file_path="交付 results/report.html")
    assert "产业链 L4-L5" in str(read), read
    edited = await invoke(
        "Edit",
        file_path="交付 results/report.html",
        old_string="产业链 L4-L5",
        new_string="产业链节点最终版",
    )
    assert edited.get("ok"), edited
    found = await invoke("Glob", pattern="*.html", path="交付 results")
    assert str(generated / "report.html") in found.get("filenames", []), found
    searched = await invoke("Grep", pattern="最终版", path="交付 results/report.html")
    assert "report.html" in str(searched), searched
    written = await invoke("Write", file_path="交付 results/说明.txt", content="项目成果")
    assert written.get("ok"), written
    pin = await invoke(
        "pin_to_workspace",
        file_paths=[
            "交付 results/" + name for name in ("report.html", "nodes.xlsx", "delivery.zip")
        ],
    )
    assert pin.get("ok"), pin
    ids = {item["name"]: item["file_id"] for item in pin["pinned"]}
    assert all(fid.startswith("lpf_") for fid in ids.values()), pin

    monkeypatch.setattr(
        "core.auth.backend.settings", replace(settings, auth=replace(settings.auth, mode="session"))
    )
    monkeypatch.setenv("HUGAGENT_DESKTOP_BRIDGE_SECRET", "cwd-e2e-bridge")
    monkeypatch.delenv("HUGAGENT_CAPS_ROOT", raising=False)
    db_session.query(UserShadow).filter_by(user_id="pin-user").one().user_center_id = (
        "cwd-preview-user"
    )
    db_session.commit()
    app = FastAPI()
    app.include_router(files.router)
    app.dependency_overrides[get_db] = lambda: db_session
    headers = {
        "x-desktop-bridge": "cwd-e2e-bridge",
        "x-desktop-bridge-user": base64.b64encode(
            json.dumps({"user_center_id": "cwd-preview-user", "username": "Preview"}).encode()
        ).decode(),
    }
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://preview"
    ) as client:
        for name, fid in ids.items():
            location = await client.get(f"/files/{fid}/local-path", headers=headers)
            assert location.status_code == 200, location.text
            assert str(generated / name) in location.text
            preview = await client.get(f"/files/{fid}?inline=1", headers=headers)
            assert preview.status_code == 200, preview.text
            if name.endswith(".html"):
                assert "产业链节点最终版" in preview.text
                (generated / name).write_text("<h1>Live update after pin</h1>")
                live = await client.get(f"/files/{fid}?inline=1", headers=headers)
                assert "Live update after pin" in live.text
            elif name.endswith(".xlsx"):
                book = load_workbook(io.BytesIO(preview.content))
                assert list(book.active.values) == [("L4", "L5"), ("材料", "高纯材料")]
                book.close()
            else:
                with zipfile.ZipFile(io.BytesIO(preview.content)) as archive:
                    assert archive.testzip() is None
                    assert set(archive.namelist()) == {"report.html", "nodes.xlsx"}
            assert (await client.get(f"/files/{fid}?inline=1")).status_code == 401
    await provider.close_session("cwd-child")
    assert all((generated / name).is_file() for name in ids)


@pytest.mark.parametrize("change", ["removed", "rebound", "deleted"])
async def test_project_changed_during_run_blocks_tools_without_fallback(
    local_tools, db_session, change
):
    from core.db.models import Project
    from datetime import datetime, timezone

    invoke, root, scope, provider = local_tools
    if change == "removed":
        root.rename(root.with_name("moved-project"))
    else:
        project = db_session.query(Project).filter_by(project_id="pin-project").one()
        if change == "rebound":
            project.extra_data = {"local": {"path": str(root.with_name("other-project"))}}
        else:
            project.deleted_at = datetime.now(timezone.utc)
        db_session.commit()
    for name, args in (
        ("Write", dict(file_path="unexpected.txt", content="bad")),
        ("bash", dict(command="printf bad > unexpected.txt")),
    ):
        result = await invoke(name, **args)
        assert result.get("blocked"), result
    assert not (root / "unexpected.txt").exists()


@pytest.mark.parametrize("output_name", ["result.txt", "123"])
async def test_relative_permissions_use_project_without_implicitly_granting_it(
    local_project, output_name
):
    from core.llm.tool_permissions import (
        ToolPermissionRegistry,
        ToolPermissionService,
        PermissionRuntime,
        builtin_tool_permission,
    )
    from core.services.local_grant_service import remove_grant

    source, scope = local_project
    registry = ToolPermissionRegistry()
    for name in ("Write", "Read", "bash"):
        registry.register(name, builtin_tool_permission(name), source="integration")
    service = ToolPermissionService(
        registry,
        PermissionRuntime(
            chat_id="pin-chat",
            sandbox_session_id="permission-child",
            user_id="pin-user",
            interactive=False,
            approval_available=False,
            approval_mode="ask",
            project_scope=scope,
        ),
    )

    async def check(name, **args):
        return await service.authorize(SimpleNamespace(name=name, id="permission-call", input=args))

    allowed = await check("Write", file_path="result.txt", content="x")
    assert allowed.proceed, allowed.payload
    assert {i.target for i in allowed.ticket.intents} == {str(source.parent / "result.txt")}
    command = await check("bash", command=f"printf x > {output_name}")
    assert command.proceed, command.payload
    assert command.ticket.local_command.cwd == str(source.parent)
    outside = source.parent.parent / "outside.txt"
    outside.write_text("private")
    (source.parent / "escape.txt").symlink_to(outside)
    assert not (await check("Read", file_path="escape.txt")).proceed
    remove_grant(str(source.parent))
    for name, args in (
        ("Read", {"file_path": "report.txt"}),
        ("Write", {"file_path": "result.txt", "content": "x"}),
        ("bash", {"command": f"printf x > {output_name}"}),
    ):
        denied = await check(name, **args)
        assert not denied.proceed, (name, denied)
        assert denied.payload.get("error"), denied.payload


async def test_confined_bash_uses_project_cwd(local_tools):
    from dataclasses import replace

    invoke, root, scope, provider = local_tools
    invoke.service.runtime = replace(invoke.service.runtime, approval_mode="ask")
    result = await invoke("bash", command="pwd; printf confined > confined.txt")
    assert result.get("exit_code") == 0, result
    assert result["stdout"].strip() == str(root)
    assert (root / "confined.txt").read_text() == "confined"


async def test_real_factory_agent_uses_bound_project_for_file_and_bash(local_tools, monkeypatch):
    from agentscope.message import UserMsg, TextBlock, ToolCallBlock
    from agentscope.model import ChatResponse
    from core.llm import agent_factory

    invoke, root, scope, provider = local_tools

    class EmptyLoader:
        def load_all_metadata(self):
            return {}

        def get_skill_dir(self, *args):
            return None

        def register_skills_to_toolkit(self, *args, **kwargs):
            return 0

    class ScriptedModel:
        model = "cwd-e2e-scripted"
        context_size = 32768
        calls = 0

        async def count_tokens(self, messages, tools=None):
            return 1

        async def __call__(self, messages, tools=None, **kwargs):
            self.calls += 1
            if self.calls == 1:
                content = [
                    ToolCallBlock(
                        id="factory-write",
                        name="Write",
                        input=json.dumps(
                            {"file_path": "factory-write.txt", "content": "factory write"}
                        ),
                    )
                ]
            elif self.calls == 2:
                content = [
                    ToolCallBlock(
                        id="factory-bash",
                        name="bash",
                        input=json.dumps({"command": "printf factory-bash > factory-bash.txt"}),
                    )
                ]
            else:
                content = [TextBlock(type="text", text="done")]
            return ChatResponse(content=content, is_last=True)

    monkeypatch.setattr(agent_factory, "get_skill_loader", lambda: EmptyLoader())
    from core.db.engine import SessionLocal
    from core.services.run_journal import RunJournal

    journal = RunJournal(SessionLocal)
    journal.accept(
        run_id="cwd-factory-run",
        message_id="cwd-factory-message",
        chat_id="pin-chat",
        user_id="pin-user",
        request_payload={"kind": "chat"},
        recovery_snapshot={"kind": "chat", "worker_args": {}},
    )
    assert journal.claim("cwd-factory-run", owner="cwd-test-worker", lease_seconds=300)
    agent, clients = await agent_factory.create_agent_executor(
        enabled_skill_ids=[],
        enabled_mcp_ids=[],
        enabled_kb_ids=[],
        max_iters=4,
        chat_id="pin-chat",
        sandbox_session_id="factory-child",
        current_user_id="pin-user",
        run_id="cwd-factory-run",
        journal_owner="cwd-test-worker",
        approval_mode="full",
        project_ctx={
            "project_id": "pin-project",
            "project_is_local": True,
            "project_local_path": str(root),
            "project_name": "Reports",
        },
    )
    agent.model = ScriptedModel()
    agent.state.model_pinned = True
    try:
        reply = await agent.reply(UserMsg(name="user", content="Write the two project files."))
        assert "done" in reply.get_text_content()
        assert (root / "factory-write.txt").read_text() == "factory write"
        assert (root / "factory-bash.txt").read_text() == "factory-bash"
        assert f"<cwd>{root}</cwd>" in agent._jx_compaction_system_prompt
    finally:
        for client in clients:
            await client.close()
        await provider.close_session("factory-child")
