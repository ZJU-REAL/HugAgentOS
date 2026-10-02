"""Actual LibreOffice cold/warm conversion, isolated synthetic DOCX/cache."""

import json
from pathlib import Path
import shutil
import statistics
import tempfile
import time
from docx import Document
from api.routes.files import _convert_office_to_pdf
from core.content.office_preview_cache import cached_office_pdf

with tempfile.TemporaryDirectory(prefix="perf172-office-") as directory:
    root = Path(directory)
    document = Document()
    for i in range(20):
        document.add_heading(f"Performance section {i}", 2)
        document.add_paragraph("Synthetic document for preview benchmarking. " * 50)
    source = root / "sample.docx"
    document.save(source)
    result = {}
    for name, count in [("uncached", 3), ("cached", 6)]:
        times = []
        for _ in range(count):
            start = time.perf_counter()
            pdf, work = (
                _convert_office_to_pdf(str(source), "performance")
                if name == "uncached"
                else cached_office_pdf(
                    str(source), "performance", _convert_office_to_pdf, root=root / "cache"
                )
            )
            assert Path(pdf).read_bytes().startswith(b"%PDF-")
            times.append(round((time.perf_counter() - start) * 1000, 2))
            shutil.rmtree(work)
        result[name] = times
    result["warm_p50_ms"] = statistics.median(result["cached"][1:])
    Path(".git/performance-172/preview-results.json").write_text(json.dumps(result, indent=2))
    print(json.dumps(result, indent=2))
